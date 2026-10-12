#!/usr/bin/env python3
"""Build the 10-slide interview deck (portfolio/presentation/interview-deck.pptx) with speaker notes.

Requires python-pptx (pip install python-pptx==1.0.2). Screenshots come from portfolio/screenshots
(demo/application/scripts/capture-screenshots.mjs).
  python3 scripts/build_interview_deck.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "portfolio/screenshots"
OUT = ROOT / "portfolio/presentation/interview-deck.pptx"

INK = RGBColor(0x17, 0x20, 0x2B)
MUTED = RGBColor(0x55, 0x61, 0x71)
ACCENT = RGBColor(0x0B, 0x5C, 0xAD)
SOFT = RGBColor(0xF1, 0xF4, 0xF8)
CRIT = RGBColor(0xA4, 0x16, 0x1A)
OK = RGBColor(0x1C, 0x6B, 0x3A)
WARN = RGBColor(0x8A, 0x4B, 0x00)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height
M = Inches(0.6)


def text(slide, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, spacing=1.1):
    """runs: str or list of paragraphs; a paragraph is str or list of (text, dict(bold/color/size))."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    paragraphs = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        p.space_after = Pt(6)
        parts = para if isinstance(para, list) else [(para, {})]
        for t, style in parts:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(style.get("size", size))
            r.font.bold = style.get("bold", bold)
            r.font.color.rgb = style.get("color", color)
            r.font.name = "Helvetica Neue"
    return box


def rect(slide, x, y, w, h, fill=SOFT):
    shape = slide.shapes.add_shape(1, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def header(slide, kicker, title):
    text(slide, M, Inches(0.35), W - 2 * M, Inches(0.4), kicker.upper(), size=12, color=MUTED)
    text(slide, M, Inches(0.7), W - 2 * M, Inches(0.9), title, size=30, bold=True)


def footer(slide, n):
    text(slide, M, H - Inches(0.45), W - 2 * M, Inches(0.3),
         f"Fictional organisation · portfolio case · {n}/10", size=10, color=MUTED)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def picture(slide, name, x, y, w=None, h=None):
    return slide.shapes.add_picture(str(SHOTS / name), x, y, width=w, height=h)


def card(slide, x, y, w, h, title, body, title_color=ACCENT, size=15):
    rect(slide, x, y, w, h)
    text(slide, x + Inches(0.2), y + Inches(0.15), w - Inches(0.4), Inches(0.5), title, size=17, bold=True, color=title_color)
    text(slide, x + Inches(0.2), y + Inches(0.65), w - Inches(0.4), h - Inches(0.8), body, size=size, color=INK)


# 1 — title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, fill=RGBColor(0x0F, 0x2A, 0x47))
text(s, M, Inches(1.6), Inches(11), Inches(0.5), "BUSINESS ANALYSIS PORTFOLIO CASE", size=14, color=RGBColor(0xA9, 0xC8, 0xEA))
text(s, M, Inches(2.1), Inches(11.5), Inches(2), "Shipment exceptions, handled before the customer calls", size=44, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
text(s, M, Inches(4.1), Inches(11), Inches(1.2), "From three conflicting requests to rules that run as tests, a working demo and a controlled pilot — for a fictional short-sea and door-to-door operator.", size=20, color=RGBColor(0xDD, 0xE7, 0xF2))
text(s, M, Inches(6.4), Inches(11), Inches(0.5), "Philippe Godfroy · fictional organisation, simulated delivery, real automated tests", size=14, color=RGBColor(0xA9, 0xC8, 0xEA))
notes(s, "Open with the one-sentence version: customers heard about delays at the gate; I turned three requested solutions into one problem and worked it out end to end. Say upfront that the organisation is fictional and the delivery simulated — and that the tests and demo are real.")

# 2 — the problem
s = prs.slides.add_slide(BLANK)
header(s, "1 · The problem", "Customers learned about delays at the gate")
stats = [("22%", "of late deliveries announced before the window", CRIT), ("38%", "of all contacts are “where is my load?”", WARN), ("±55 / week", "missed delivery slots at ±€140", WARN)]
for i, (num, label, col) in enumerate(stats):
    x = M + i * Inches(4.1)
    rect(s, x, Inches(1.8), Inches(3.8), Inches(1.7))
    text(s, x + Inches(0.25), Inches(1.95), Inches(3.4), Inches(0.8), num, size=36, bold=True, color=col)
    text(s, x + Inches(0.25), Inches(2.75), Inches(3.4), Inches(0.7), label, size=15, color=MUTED)
text(s, M, Inches(3.9), Inches(12), Inches(0.5), "Three managers, three solutions:", size=20, bold=True)
text(s, M, Inches(4.45), Inches(12), Inches(1.6), [
    "“A tracking page like the parcel carriers”  —  Head of Customer Service",
    "“Proactive e-mails to key accounts”  —  Commercial director",
    "“Nothing that creates more work”  —  Road planning lead",
], size=18, color=INK)
text(s, M, Inches(6.2), Inches(12), Inches(0.5), [[("Root cause: ", {"bold": True, "color": ACCENT}), ("no shared definition of late, no severity, no owner for telling the customer.", {})]], size=18)
footer(s, 2)
notes(s, "Illustrative figures. The point of the slide: the request came in as three solutions. My first job was to describe the problem without choosing one of them. The 5 whys ended at 'missing rules', not 'missing screen'.")

# 3 — approach
s = prs.slides.add_slide(BLANK)
header(s, "2 · Approach", "From opportunity to production readiness")
steps = ["Frame & business case", "Discover & event storm", "Model AS-IS / TO-BE", "Rules & examples", "Prioritise with the PO", "Design & demo", "BAT & go/no-go", "Measure"]
for i, st in enumerate(steps):
    col, row = i % 4, i // 4
    x = M + col * Inches(3.05)
    y = Inches(1.9) + row * Inches(1.55)
    rect(s, x, y, Inches(2.85), Inches(1.3), fill=SOFT)
    text(s, x + Inches(0.2), y + Inches(0.15), Inches(0.6), Inches(0.5), str(i + 1), size=22, bold=True, color=ACCENT)
    text(s, x + Inches(0.75), y + Inches(0.2), Inches(2.0), Inches(1.0), st, size=17, bold=True)
text(s, M, Inches(5.3), Inches(12), Inches(1.4), [
    [("BA prepares, owners decide. ", {"bold": True}), ("Scope and priority: Product Owner. Thresholds: PO with the planning lead. Go-live: PO. Funding and stage gate: sponsor.", {})],
    [("Evidence: ", {"bold": True}), ("documents with IDs that trace from pain point to KPI; a script checks every reference in CI.", {})],
], size=17)
footer(s, 3)
notes(s, "Walk the journey in 30 seconds. Emphasise the split between preparing and deciding — it comes back in every step and answers 'aren't you doing the PO's job?' before it is asked.")

# 4 — framing
s = prs.slides.add_slide(BLANK)
header(s, "3 · Framing", "Four options, one honest business case")
rows = [("O1 Add a CS agent", "2.45", "treats the symptom"), ("O2 Standard TMS e-mails", "3.10", "±9 changes per order, ±0.3 relevant"), ("O3 Exception management", "3.95", "recommended — as a pilot"), ("O4 Visibility platform", "2.85", "later, for consignees")]
tbl = s.shapes.add_table(len(rows) + 1, 3, M, Inches(1.8), Inches(7.2), Inches(2.6)).table
for j, h in enumerate(["Option", "Score", "Verdict"]):
    tbl.cell(0, j).text = h
for i, r in enumerate(rows, start=1):
    for j, v in enumerate(r):
        tbl.cell(i, j).text = v
for row in tbl.rows:
    for c in row.cells:
        for p in c.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(15)
tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(3.0), Inches(1.0), Inches(3.2)
card(s, Inches(8.2), Inches(1.8), Inches(4.5), Inches(4.3), "Riskiest assumption",
     ["The case pays back in ±13.5 months if a third of WISMO contacts disappear. Nobody knew that number.",
      "Pessimistic case: ±31 months.",
      "So the sponsor funded a six-week pilot with a control group, and release 2 only if WISMO falls ≥ 20%."], size=15)
text(s, M, Inches(4.7), Inches(7.2), Inches(1.5), "Weights were agreed before scoring, with the PO, the Head of CS and the planning lead. Proposed figures — not measured.", size=15, color=MUTED)
footer(s, 4)
notes(s, "Why not O2: I tested it on paper with a week of status events for 20 orders. More messages is not better information. The strongest move here is 'buy the information first': name the riskiest assumption and test it before the full investment.")

# 5 — event storming
s = prs.slides.add_slide(BLANK)
header(s, "4 · Workshop", "One disagreement, one missing concept: severity")
card(s, M, Inches(1.8), Inches(3.8), Inches(2.4), "Account manager", "“Key accounts must hear everything.”\nNeed: never surprised by something serious.", title_color=WARN)
card(s, Inches(4.75), Inches(1.8), Inches(3.8), Inches(2.4), "Senior planner", "“No notification storm.”\nNeed: no message that doesn't change the customer's day.", title_color=WARN)
card(s, Inches(8.9), Inches(1.8), Inches(3.8), Inches(2.4), "Decision (PO, D-02)", "By severity: minor < 2 h, major 2–6 h, critical ≥ 6 h or next day; reefer +1; critical key account also gets a call.", title_color=OK)
text(s, M, Inches(4.6), Inches(12), Inches(1.8), [
    [("How: ", {"bold": True}), ("ask each side for the last load that went wrong → reflect the need back → put both needs on the wall → let the room find the concept → three options → the owner decides.", {})],
    [("Honest note: ", {"bold": True}), ("the workshop is simulated. The facilitation plan and choices are real; I have not yet run event storming with a real group.", {})],
], size=17)
footer(s, 5)
notes(s, "This is the stakeholder-conflict story. Tell it with the dialogue in docs/03-workshops/workshop-dialogue.md. Severity resolved three hotspots at once. Keep the honest note — interviewers respect it.")

# 6 — rules as tests
s = prs.slides.add_slide(BLANK)
header(s, "5 · Requirements", "Rules precise enough to run as tests")
picture(s, "07-acceptance-examples-and-tests.png", M, Inches(1.7), w=Inches(7.4))
text(s, Inches(8.3), Inches(1.8), Inches(4.4), Inches(4.8), [
    [("13 rules, one decision table, 30 Given/When/Then examples", {"bold": True})],
    "The examples are the specification and the test: the prototype, the demo and CI all use the same rule module.",
    "When the PO changes a threshold, the failing examples show exactly which decisions change meaning.",
    [("Traceability: ", {"bold": True}), ("P5 → BR-1 → FR-06/10 → ER-10/12 → US-05/07 → AC-07.1 → K1, K3", {})],
], size=16)
footer(s, 6)
notes(s, "Boundaries matter: 119 minutes is minor, 120 is major. Explain why the window check is the first row of the table: a delayed sailing must not alarm customers whose delivery is still on time. The analyst does not write production code — the module proves the rules are precise enough to build and test.")

# 7 — backlog
s = prs.slides.add_slide(BLANK)
header(s, "6 · Product delivery", "Prioritised with the Product Owner")
rows = [("EN-2 Window mandatory for pilot customers", "15.0"), ("US-07 No message storm", "9.0"), ("US-13 Pilot toggle + measurement", "8.0"), ("US-05 Proactive customer message", "5.8"), ("US-06 EDI/API status messages → R2", "1.8")]
tbl = s.shapes.add_table(len(rows) + 1, 2, M, Inches(1.8), Inches(7.0), Inches(3.0)).table
tbl.cell(0, 0).text, tbl.cell(0, 1).text = "Backlog item", "WSJF"
for i, (a, b) in enumerate(rows, start=1):
    tbl.cell(i, 0).text, tbl.cell(i, 1).text = a, b
for row in tbl.rows:
    for c in row.cells:
        for p in c.text_frame.paragraphs:
            for run in p.runs:
                run.font.size = Pt(15)
tbl.columns[0].width, tbl.columns[1].width = Inches(5.8), Inches(1.2)
card(s, Inches(8.1), Inches(1.8), Inches(4.6), Inches(4.4), "Trade-offs the PO made",
     ["Missing-status hold before EDI messages: false alarms in week 1 would kill trust.",
      "Key-account call task kept: cheap, and it wins the account managers.",
      "Workbench scope cut: filters to R2.",
      "Release 1 = 36 points, two sprints, five pilot customers."], size=15)
text(s, M, Inches(5.1), Inches(7.0), Inches(1.2), "WSJF ranks; dependencies sequence. MoSCoW says what is in the release; WSJF orders the work inside it.", size=15, color=MUTED)
footer(s, 7)
notes(s, "When everything is 'Must', WSJF made the 1-point window fix rise to the top: 28% of orders had no window, and nothing works without it. Business scores value, the team scores size, the PO decides.")

# 8 — demo
s = prs.slides.add_slide(BLANK)
header(s, "7 · Design & demo", "The control tower: the release, working")
picture(s, "03-exception-critical.png", M, Inches(1.7), w=Inches(6.0))
picture(s, "04-customer-view-delayed.png", Inches(6.75), Inches(1.7), w=Inches(6.0))
text(s, M, Inches(5.65), Inches(12.2), Inches(1.2), [
    [("Planner: ", {"bold": True}), ("most urgent first; every exception explains which rules decided it.   ", {}), ("Customer: ", {"bold": True}), ("one clear message — nothing internal.", {})],
    [("Building it was a review: ", {"bold": True}), ("five open questions the spec had missed, and a privacy leak in the first customer API — found and fixed with a test.", {})],
], size=15)
footer(s, 8)
notes(s, "If there is time, open the live demo and run the guided scenario (±3 minutes): publish the sailing delay, show that 4 of 11 orders become exceptions, the critical classification, assign, act, customer view, new window, audit trail, test results.")

# 9 — BAT and go-live
s = prs.slides.add_slide(BLANK)
header(s, "8 · Acceptance & go-live", "Accepted by the business — and my own mistake")
card(s, M, Inches(1.8), Inches(3.9), Inches(4.2), "BAT (simulated)", ["Key users incl. the sceptical senior planner.", "3 of 4 majors came from real readers in their own context — the team's tests had passed.", "Signed off with 2 known issues, each with workaround, owner and date."], size=15)
card(s, Inches(4.75), Inches(1.8), Inches(3.9), Inches(4.2), "My mistake (DEF-04)", ["In refinement I answered: throttle per exception.", "In BAT a customer got two e-mails in ten minutes about one order.", "Fix: per order. Retro action: customer-facing questions go to the PO the same day."], title_color=CRIT, size=15)
card(s, Inches(8.9), Inches(1.8), Inches(3.8), Inches(4.2), "Ready for production", ["13 go/no-go criteria, an owner each.", "Shadow mode, staged switch-on on a Monday, rollback in 4 minutes.", "Support article: “rule or bug?”."], title_color=OK, size=15)
footer(s, 9)
notes(s, "DEF-04 is the best story in the case: my own analysis decision, found by a real reader, fixed, and turned into a team practice. Tell it plainly.")

# 10 — claims and questions
s = prs.slides.add_slide(BLANK)
header(s, "9 · What this shows", "Shown here — and not claimed")
card(s, M, Inches(1.8), Inches(5.9), Inches(3.6), "Shown", ["How I frame a problem, prepare decisions and specify precisely.", "Artefacts a team can use: BPMN, rules, stories, API contract, BAT, go/no-go.", "Real automated tests and a working demo.", "A learning lab to defend every artefact."], title_color=OK, size=15)
card(s, Inches(6.8), Inches(1.8), Inches(5.9), Inches(3.6), "Not claimed", ["Real stakeholders, sprints or BAT — all simulated.", "Years in a professional Scrum team.", "That the case was built without help: it was built with AI assistance; what I reviewed is tracked openly."], title_color=CRIT, size=15)
text(s, M, Inches(5.7), Inches(12), Inches(1.0), [[("My questions for you: ", {"bold": True, "color": ACCENT}), ("How does refinement work in your team, and who writes the acceptance criteria? Where do business analysts add the most value here today?", {})]], size=17)
footer(s, 10)
notes(s, "Close by being precise about evidence. Then switch to their world with the two questions. Only claim AI-assisted work at the level recorded in learning/ai-provenance.md.")

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f"Wrote {OUT.relative_to(ROOT)} ({len(prs.slides)} slides)")
